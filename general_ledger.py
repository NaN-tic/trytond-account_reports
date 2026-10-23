
# This file is part of account_reports for tryton.  The COPYRIGHT file
# at the top level of this repository contains the full copyright notices and
# license terms.
from datetime import timedelta, datetime
from decimal import Decimal
from pathlib import Path

from fpdf import FPDF
from trytond.pool import Pool, PoolMeta
from trytond.transaction import Transaction
from trytond.model import ModelView, fields
from trytond.wizard import Wizard, StateView, StateReport, Button
from trytond.pyson import Eval, Bool, If
from trytond.tools import grouped_slice, slugify
from trytond.report import Report
from trytond.i18n import gettext
from trytond.exceptions import UserError
from trytond.modules.account_reports.common import (
    AccountReportsReportMixin, TimeoutException, TimeoutChecker)
from trytond.modules.account_reports.tools import vat_label
from trytond.modules.account_reports.xlsx import (
    XlsxReport, save_workbook, convert_str_to_float)
from trytond.modules.html_report.engine import (
    HTMLReportMixin, render as html_render)
from trytond.modules.xgettext import _

from trytond.rpc import RPC
from trytond.modules.account.exceptions import FiscalYearNotFoundError
from openpyxl import Workbook
_ZERO = Decimal(0)

class PrintGeneralLedgerStart(ModelView):
    'Print General Ledger'
    __name__ = 'account_reports.print_general_ledger.start'
    fiscalyear = fields.Many2One('account.fiscalyear', 'Fiscal Year',
        states={
            'invisible': Eval('start_date') | Eval('end_date'),
            'required': ~Eval('start_date') & ~Eval('end_date'),
            })
    start_period = fields.Many2One('account.period', 'Start Period',
        states={
            'invisible': Eval('start_date') | Eval('end_date'),
            'required': ~Eval('start_date') & ~Eval('end_date'),
            },
        domain=[
            ('fiscalyear', '=', Eval('fiscalyear')),
            If(Bool(Eval('end_period')),
                ('start_date', '<=', (Eval('end_period'), 'start_date')),
                (),
                ),
            ])
    end_period = fields.Many2One('account.period', 'End Period',
        states={
            'invisible': Eval('start_date') | Eval('end_date'),
            'required': ~Eval('start_date') & ~Eval('end_date'),
            },
        domain=[
            ('fiscalyear', '=', Eval('fiscalyear')),
            If(Bool(Eval('start_period')),
                ('start_date', '>=', (Eval('start_period'), 'start_date')),
                (),
                ),
            ])
    start_date = fields.Date('Initial Posting Date',
        domain=[
            If(Eval('start_date') & Eval('end_date'),
                ('start_date', '<=', Eval('end_date', None)),
                ()),
            ],
        states={
            'invisible': Eval('start_period') | Eval('end_period'),
            'required': ((Eval('start_date') | Eval('end_date')) &
                ~Bool(Eval('start_period') | Eval('end_period'))),
            })
    end_date = fields.Date('Final Posting Date',
        domain=[
            If(Eval('start_date') & Eval('end_date'),
                ('end_date', '>=', Eval('start_date', None)),
                ()),
            ],
        states={
            'invisible': Eval('start_period') | Eval('end_period'),
            'required': ((Eval('end_date') | Eval('start_date')) &
                ~Bool(Eval('start_period') | Eval('end_period')))
            })
    accounts = fields.Many2Many('account.account', None, None, 'Accounts')
    all_accounts = fields.Boolean('All accounts with and without balance',
        states={
            'readonly': Bool(Eval('accounts')),
            },
        help='If unchecked only print accounts with previous balance different'
        ' from 0 or with moves')
    parties = fields.Many2Many('party.party', None, None, 'Parties',
        context={
            'company': Eval('company', -1),
            },
        depends=['company'])
    output_format = fields.Selection([
            ('pdf', 'PDF'),
            ('xlsx', 'Excel'),
            ], 'Output Format', required=True)
    company = fields.Many2One('company.company', 'Company', required=True)
    timeout = fields.Integer('Timeout (s)', required=True, help='If report '
        'calculation should take more than the specified timeout (in seconds) '
        'the process will be stopped automatically.')
    show_description = fields.Boolean('Show Description',
        help='If checked show description from Account Move Line')

    @staticmethod
    def default_fiscalyear():
        pool = Pool()
        FiscalYear = pool.get('account.fiscalyear')
        try:
            fiscalyear = FiscalYear.find(
                Transaction().context.get('company'), test_state=False)
        except FiscalYearNotFoundError:
            return None
        return fiscalyear.id

    @staticmethod
    def default_all_accounts():
        return True

    @staticmethod
    def default_company():
        return Transaction().context.get('company')

    @staticmethod
    def default_output_format():
        return 'pdf'

    @staticmethod
    def default_timeout():
        Config = Pool().get('account.configuration')
        config = Config(1)
        return config.default_timeout or 60

    def default_show_description():
        return False

    @fields.depends('fiscalyear')
    def on_change_fiscalyear(self):
        self.start_period = None
        self.end_period = None

    @fields.depends('accounts')
    def on_change_accounts(self):
        if self.accounts:
            self.all_accounts = False

class PrintGeneralLedger(Wizard):
    'Print General Ledger'
    __name__ = 'account_reports.print_general_ledger'
    start = StateView('account_reports.print_general_ledger.start',
        'account_reports.print_general_ledger_start_view_form', [
            Button('Cancel', 'end', 'tryton-cancel'),
            Button('Print', 'print_', 'tryton-print', default=True),
            ])
    print_ = StateReport('account_reports.general_ledger')

    def do_print_(self, action):
        start_period = None
        if self.start.start_period:
            start_period = self.start.start_period.id
        end_period = None
        if self.start.end_period:
            end_period = self.start.end_period.id
        data = {
            'company': self.start.company.id,
            'fiscalyear': self.start.fiscalyear.id if self.start.fiscalyear else None,
            'start_period': start_period,
            'end_period': end_period,
            'start_date': self.start.start_date,
            'end_date': self.start.end_date,
            'accounts': [x.id for x in self.start.accounts],
            'all_accounts': self.start.all_accounts,
            'parties': [x.id for x in self.start.parties],
            'output_format': self.start.output_format,
            'timeout': self.start.timeout,
            'show_description': self.start.show_description,
            }
        if self.start.output_format == 'xlsx':
            ActionReport = Pool().get('ir.action.report')
            action_report, = ActionReport.search([
                    ('report_name', '=', 'account_reports.general_ledger_xlsx'),
                    ])
            action = action_report.action.get_action_value()
        return action, data

    def transition_print_(self):
        return 'end'

    def default_start(self, fields):
        Party = Pool().get('party.party')
        account_ids = []
        party_ids = []
        if Transaction().context.get('model') == 'party.party':
            for party in Party.browse(Transaction().context.get('active_ids')):
                if party.account_payable:
                    account_ids.append(party.account_payable.id)
                if party.account_receivable:
                    account_ids.append(party.account_receivable.id)
                party_ids.append(party.id)
        return {
            'accounts': account_ids,
            'parties': party_ids,
            }


class GeneralLedgerReport(AccountReportsReportMixin, HTMLReportMixin, Report):
    __name__ = 'account_reports.general_ledger'

    @classmethod
    def __setup__(cls):
        super(GeneralLedgerReport, cls).__setup__()
        cls.__rpc__['execute'] = RPC(False)
        cls.side_margin = 0.3

    @classmethod
    def _ref_origin_invoice_line(cls, line):
        if not line.origin or not line.origin.invoice:
            return ''
        ref = []
        if line.origin.invoice.number:
            ref.append('%s' % line.origin.invoice.number)
        if line.origin.invoice.reference:
            ref.append('[%s]' % line.origin.invoice.reference)
        if line.origin.invoice.party.accounting_name:
            ref.append('%s' % line.origin.invoice.party.accounting_name)
        return ' '.join(ref)

    @classmethod
    def _ref_origin_invoice(cls, line):
        ref = []
        if line.move_origin.number:
            ref.append('%s' % line.move_origin.number)
        if line.move_origin.reference:
            ref.append('[%s]' % line.move_origin.reference)
        if line.move_origin.party.accounting_name:
            ref.append('%s' % line.move_origin.party.accounting_name)
        return ' '.join(ref)

    @classmethod
    def _ref_origin_bank_line(cls, line):
        if line.origin.description:
            ref = '%s' % line.origin.description
        else:
            ref = (line.origin.rec_name if line.origin
                and hasattr(line.origin, 'rec_name') else None)
        return ref

    @classmethod
    def _ref_origin_statement_origin(cls, line):
        if line.origin.description and not line.description:
            ref = '%s' % line.origin.description
        else:
            ref = (line.origin.rec_name if line.origin
                and hasattr(line.origin, 'rec_name') else None)
        return ref

    @classmethod
    def _ref_origin(cls, line):
        return (line.origin.rec_name if line.origin
            and hasattr(line.origin, 'rec_name') else None)

    @classmethod
    def _ref(cls, line):
        ref = line.description_used or line.move_description_used or None
        return (ref if ref else (line.move_origin.rec_name if line.move_origin
                and hasattr(line.move_origin, 'rec_name') else None))

    @classmethod
    def prepare(cls, data, checker):
        pool = Pool()
        Company = pool.get('company.company')
        FiscalYear = pool.get('account.fiscalyear')
        Period = pool.get('account.period')
        Account = pool.get('account.account')
        Party = pool.get('party.party')
        Line = pool.get('account.move.line')
        Invoice = pool.get('account.invoice')
        InvoiceLine = pool.get('account.invoice.line')
        try:
            Origin = pool.get('account.bank.statement.line')
        except:
            Origin = None
        try:
            BankLine = pool.get('account.bank.statement.line')
        except:
            BankLine = None

        def _get_key(currentKey):
            party = (currentKey[1].accounting_name if len(currentKey) > 1
                and currentKey[1] else 'None')
            return (currentKey[0].code, party or 'None')

        fiscalyear = (FiscalYear(data['fiscalyear']) if data.get('fiscalyear')
            else None)
        start_period = None
        if data['start_period']:
            start_period = Period(data['start_period'])
        end_period = None
        if data['end_period']:
            end_period = Period(data['end_period'])
        start_date = data.get('start_date', None)
        end_date = data.get('end_date', None)
        with Transaction().set_context(active_test=False):
            accounts = Account.browse(data.get('accounts', []))
            parties = Party.browse(data.get('parties', []))
            if accounts:
                accounts_subtitle = []
                for x in accounts:
                    if len(accounts_subtitle) > 4:
                        accounts_subtitle.append('...')
                        break
                    accounts_subtitle.append(x.code)
                accounts_subtitle = ', '.join(accounts_subtitle)
            else:
                accounts_subtitle = ''

            if parties:
                parties_subtitle = []
                for x in parties:
                    if len(parties_subtitle) > 4:
                        parties_subtitle.append('...')
                        break
                    parties_subtitle.append(x.accounting_name)
                parties_subtitle = '; '.join(parties_subtitle)
            else:
                parties_subtitle = ''

        if data['company']:
            company = Company(data['company'])
        elif fiscalyear:
            company = fiscalyear.company
        else:
            company = Company(Transaction().context.get('company', -1))

        parameters = {}
        parameters['company'] = company.rec_name
        parameters['company_vat'] = (company.party.tax_identifier
            and company.party.tax_identifier.code) or ''
        parameters['company_vat_label'] = (company.party.tax_identifier
            and vat_label(company.party.tax_identifier) or '')
        parameters['start_period'] = start_period and start_period or ''
        parameters['end_period'] = end_period and end_period or ''
        parameters['start_date'] = (start_date.strftime('%d/%m/%Y')
            if start_date else '')
        parameters['end_date'] = (end_date.strftime('%d/%m/%Y')
            if end_date else '')
        parameters['fiscal_year'] = fiscalyear.rec_name if fiscalyear else ''
        parameters['accounts'] = accounts_subtitle
        parameters['parties'] = parties_subtitle
        parameters['show_description'] = data.get('show_description', True)

        where = ''
        if accounts:
            where += "aml.account in (%s) " % (
                ",".join([str(a.id) for a in accounts]))
        else:
            where += "aa.parent is not null "

        if start_date:
            where += "and am.company = %s " % company.id
            where += "and am.date >= '%s' " % start_date
            where += "and am.date <= '%s' " % end_date
        else:
            filter_periods = fiscalyear.get_periods(start_period, end_period)
            where += "and am.period in (%s) " % (
                ",".join([str(a.id) for a in filter_periods]))

        if parties:
            where += " and aml.party in (%s)" % (
                ",".join([str(a.id) for a in parties]))

        cursor = Transaction().connection.cursor()
        cursor.execute("""
            SELECT
                aml.id
            FROM
                account_move_line aml,
                account_move am,
                account_account aa,
                account_account_type aat
            WHERE
                am.id = aml.move AND
                aa.id = aml.account AND
                aa.type = aat.id AND
                %s
            ORDER BY
                aml.account,
                -- Sort by party only when account is of
                -- type 'receivable' or 'payable'
                -- or party_requierd is True
                CASE WHEN aat.receivable or aat.payable or
                    aa.party_required THEN aml.party ELSE 0 END,
                am.date,
                am.id,
                am.description,
                aml.id
            """ % where)
        line_ids = [x[0] for x in cursor.fetchall()]

        if not start_date:
            start_date = (start_period.start_date if start_period
                else fiscalyear.start_date)
        initial_balance_date = start_date - timedelta(days=1)
        with Transaction().set_context(date=initial_balance_date):
            init_values = {}
            if not parties:
                init_values = Account.html_read_account_vals(accounts, company,
                    with_moves=False, exclude_party_moves=True)
            init_party_values = Party.html_get_account_values_by_party(
                parties, accounts, company)
            init_parties = set([p for a, av in init_party_values.items()
                    for p, pv in av.items() if p is not None])
        records = {}
        parties_general_ledger = set() # (account_id, party_id)
        lastKey = None
        sequence = 0
        accounts_w_moves = []
        balance = _ZERO
        # Add the asked period/date lines in records
        for group_lines in grouped_slice(line_ids):
            checker.check()
            init_balance = _ZERO
            init_party_balance = _ZERO
            for line in Line.browse(group_lines):
                if line.account not in accounts_w_moves:
                    accounts_w_moves.append(line.account.id)
                currentKey = (line.account, line.party)
                if lastKey != currentKey:
                    balance = _ZERO
                    lastKey = currentKey
                    account_id = currentKey[0].id
                    party_id = (currentKey[1].id if len(currentKey) > 1
                        and currentKey[1] else None)
                    parties_general_ledger.add((account_id, party_id))
                    init_balance = init_values.get(line.account.id,
                        {}).get('balance', _ZERO)
                    init_party_balance = init_party_values.get(account_id,
                        {}).get(party_id, {}).get('balance', _ZERO)

                credit = line.credit
                debit = line.debit
                balance += line.debit - line.credit
                if init_party_balance:
                    balance += init_party_balance
                    init_party_balance = _ZERO
                elif init_balance:
                    balance += init_balance
                    init_balance = _ZERO
                sequence += 1

                party = None
                ref = None

                if line.origin and isinstance(line.origin, InvoiceLine):
                    ref = cls._ref_origin_invoice_line(line)

                    # If the account have the check "party_required", try to
                    # get from the invoice
                    if line.account.party_required:
                        party = line.origin.invoice.party
                elif (line.move_origin
                        and isinstance(line.move_origin, Invoice)):
                    ref = cls._ref_origin_invoice(line)

                    # If the account have the check "party_required", try to
                    # get from the invoice
                    if line.account.party_required:
                        party = line.move_origin.party
                elif (line.origin and Origin
                        and isinstance(line.origin, Origin)):
                    ref = cls._ref_origin_bank_line(line)
                elif (line.origin and BankLine
                        and isinstance(line.origin, BankLine)):
                    ref = cls._ref_origin_bank_line(line)
                elif line.origin:
                    ref = cls._ref_origin(line)
                if not ref:
                    ref = cls._ref(line)

                # If we don't fill the party in a party_required account, try
                # get the party field in the line
                if line.account.party_required and not party:
                    party = line.party

                rline = {
                    'sequence': sequence,
                    'line': line,
                    'ref': ref,
                    'credit': credit,
                    'debit': debit,
                    'balance': balance,
                    'party': party
                    }

                key = _get_key(currentKey)
                if records.get(key):
                    records[key]['lines'].append(rline)
                    records[key]['total_debit'] += debit
                    records[key]['total_credit'] += credit
                else:
                    previous_balance = init_values.get(
                        line.account.id, {}).get('balance', _ZERO)
                    if not previous_balance and line.party:
                        previous_balance = init_party_values.get(
                            line.account.id, {}).get(line.party.id, {}).get(
                                'balance', _ZERO)
                    records[key] = {
                        'account': line.account.name,
                        'code': line.account.code or str(line.account.id),
                        'party': line.party.accounting_name if line.party else None,
                        'party_required': line.account.party_required,
                        'lines': [rline],
                        'previous_balance': previous_balance,
                        'total_debit': debit,
                        'total_credit': credit,
                        }

        # Control if there are some party moves with initial value, but not
        # values in the current period control moves and must be to set.
        missing_init_parties = list(
            set(init_parties) - set([p[1] for p in parties_general_ledger]))
        if missing_init_parties:
            account_ids = [k for k, _ in init_party_values.items()]
            accounts = dict((a.id, a) for a in Account.browse(account_ids))
            for k, v in init_party_values.items():
                account = accounts[k]
                for p, z in v.items():
                    if p not in missing_init_parties:
                        continue
                    party = Party(p) if p else None
                    if not party:
                        continue
                    currentKey = (account, party)
                    credit = z.get('credit', _ZERO)
                    debit = z.get('debit', _ZERO)
                    balance = z.get('balance', _ZERO)
                    if balance == _ZERO:
                        continue
                    sequence += 1
                    rline = {
                        'sequence': sequence,
                        'line': None,
                        'ref': None,
                        'credit': credit,
                        'debit': debit,
                        'balance': balance,
                        'party': party
                        }
                    key = _get_key(currentKey)
                    if records.get(key):
                        records[key]['lines'].append(rline)
                        records[key]['total_debit'] += debit
                        records[key]['total_credit'] += credit
                    else:
                        records[key] = {
                            'account': account.name,
                            'code': account.code or str(account.id),
                            'party': party.accounting_name if party else None,
                            'party_required': account.party_required,
                            'lines': [rline],
                            'previous_balance': (balance + credit - debit),
                            'total_debit': debit,
                            'total_credit': credit,
                            }
            checker.check()
        if data.get('all_accounts', True):
            # When all accounts is checked True, meaning that all accounts
            # with and without inital blaance or balance need to be printed.
            # Control if there is a missing account move in the init_values
            # list.
            init_values_account_wo_moves = {
                k: init_values[k] for k in init_values
                if k not in accounts_w_moves}
            for account_id, values in init_values_account_wo_moves.items():
                account = Account(account_id)
                balance = values.get('balance', _ZERO)
                credit = values.get('credit', _ZERO)
                debit = values.get('debit', _ZERO)
                if balance == 0:
                    continue

                currentKey = (account,)
                key = _get_key(currentKey)
                if records.get(key):
                    records[key]['total_debit'] += debit
                    records[key]['total_credit'] += credit
                else:
                    records[key] = {
                        'account': account.name,
                        'code': account.code or str(account.id),
                        'party': None,
                        'party_required': account.party_required,
                        'lines': [],
                        'previous_balance': (balance + credit - debit),
                        'total_debit': debit,
                        'total_credit': credit,
                        }
            checker.check()
            # When all accounts is checked True, meaning that all accounts
            # with and without initla blaance or blaane need to be printed.
            # Control if there is a missing account move in the
            # init_party_values list.
            account_ids = [k for k, _ in init_party_values.items()]
            accounts = dict((a.id, a) for a in Account.browse(account_ids))

            for k, v in init_party_values.items():
                account = accounts[k]
                for p, z in v.items():
                    # check if (account, party) is in current general ledger
                    if (k, p) in parties_general_ledger:
                        continue
                    balance = z.get('balance', _ZERO)
                    if balance == 0:
                        continue
                    party = Party(p)
                    currentKey = (account, party)
                    sequence += 1
                    credit = z.get('credit', _ZERO)
                    debit = z.get('debit', _ZERO)

                    key = _get_key(currentKey)
                    if not records.get(key):
                        records[key] = {
                            'account': account.name,
                            'code': account.code or str(account.id),
                            'lines': [],
                            'party': party.accounting_name if party else None,
                            'party_required': account.party_required,
                            'previous_balance': (balance + credit - debit),
                            'total_debit': debit,
                            'total_credit': credit,
                            }
            checker.check()

        return dict(sorted(records.items())), parameters

    @classmethod
    def timeout_exception(cls):
        raise TimeoutException

    @classmethod
    def execute(cls, ids, data):
        pool = Pool()
        Config = pool.get('account.configuration')
        Lang = pool.get('ir.lang')
        action, model = cls.get_action(data)
        cls.check_access(action, model or 'account.account', ids)
        timeout = data.get('timeout') or Config(1).default_timeout or 300
        checker = TimeoutChecker(timeout, cls.timeout_exception)
        side_margin = action.html_side_margin
        if side_margin is None:
            side_margin = cls.side_margin
        with Transaction().set_context(active_test=False,
                html_report_language=Lang.get(), output_format='pdf'):
            try:
                records, parameters = cls.prepare(data, checker)
                content = cls._build_pdf({
                    'records': records, 'parameters': parameters,
                    }, action, side_margin, checker)
                checker.check()
            except TimeoutException:
                raise UserError(gettext('account_reports.msg_timeout_exception'))
        if action.html_copies and action.html_copies > 1:
            content = cls.merge_pdfs([content] * action.html_copies)
        filename = slugify(cls.get_name(action))
        try:
            Printer = pool.get('printer')
        except KeyError:
            Printer = None
        if Printer:
            return Printer.send_report('pdf', content, filename, action)
        return 'pdf', content, cls.get_direct_print(action), filename

    @classmethod
    def _build_pdf(cls, data, action, side_margin, checker):
        p = data['parameters']
        filters = ['%s: %s' % (p['company_vat_label'], p['company_vat'])]
        if p['start_date']:
            filters.append(_('Start Date: %s End Date: %s') % (
                p['start_date'], p['end_date']))
        else:
            filters.append(
                _('Fiscal Year: %s Start Period: %s End Period: %s') % (
                    p['fiscal_year'],
                    p['start_period'].name if p['start_period'] else '',
                    p['end_period'].name if p['end_period'] else ''))
        filters.extend([
            _('Parties: %s') % p['parties'] if p['parties'] else _('All Parties'),
            _('Accounts: %s') % p['accounts']
                if p['accounts'] else _('All Accounts'),
            _("When move number is between parentheses it means that "
                "it has no post number and the number shown is the "
                "provisional one."),
            ])
        margin = cls.get_page_margin(action)
        # The existing report configuration stores additional margins in px.
        margin = float(side_margin) * 10 + (margin or 0) * 25.4 / 96
        pdf = LedgerPDF(
            '%s - %s - %s' % (
                p['company'], _('General Ledger'), html_render(datetime.now())),
            filters, [_('Date'), _('Number'), _('Reference // Description'),
                _('Debit'), _('Credit'), _('Balance')], margin, checker)
        pdf.set_title('%s - %s - %s' % (
            _('General Ledger'), p['company'], html_render(datetime.now())))
        pdf.add_page()
        widths = pdf.widths
        group_widths = [sum(widths[:2]), widths[2],
            sum(widths[3:5]), widths[5]]
        show_description = p.get('show_description', True)
        for record in data['records'].values():
            pdf.row([record['code'], record['party'] or record['account'],
                    _('Previous balance...') if record['lines'] else '',
                    html_render(record['previous_balance'])
                        if record['lines'] else ''],
                bold=True, fill=True, widths=group_widths,
                aligns=['L', 'L', 'R', 'R'])
            if record['lines']:
                for info in record['lines']:
                    line = info['line']
                    if line:
                        number = (line.move.number if line.move
                            and line.move.number else
                            (line.party.accounting_name if line.party else ''))
                        description = info['ref'] or ''
                        detail = (line.description
                            or line.move_description_used)
                        if show_description and detail:
                            if description:
                                description += ' // '
                            description += ' %s ' % detail
                        values = [html_render(line.date), number, description]
                    else:
                        values = ['', '-', _('Previous balance')]
                    pdf.row(values + [html_render(info[key])
                        for key in ('debit', 'credit', 'balance')])
            else:
                pdf.row(['', '-', _('Previous balance'),
                    html_render(record['total_debit']),
                    html_render(record['total_credit']),
                    html_render(record['total_debit'] - record['total_credit'])])
            pdf.row([_('Total Fiscal Year'),
                    html_render(record['total_debit']),
                    html_render(record['total_credit']),
                    html_render(record['total_debit'] - record['total_credit'])],
                bold=True, widths=[sum(widths[:3])] + widths[3:],
                aligns=['R'] * 4)
            pdf.row([record['code'], record['party'] or record['account'],
                    _('Total'), html_render(record['previous_balance']
                        + record['total_debit'] - record['total_credit'])],
                bold=True, widths=group_widths, aligns=['L', 'L', 'L', 'R'],
                bottom=True)
        return bytes(pdf.output())


class LedgerPDF(FPDF):
    """Stream ledger rows, including rows taller than a page."""

    def __init__(self, heading, filters, columns, margin, checker):
        super().__init__(orientation='L', format='A4')
        self.heading = heading
        self.filters = filters
        self.columns = columns
        self.checker = checker
        self.set_margins(margin, margin, margin)
        self.set_auto_page_break(False, margin=max(10, margin))
        self.widths = [self.epw * ratio for ratio in
            (0.08, 0.12, 0.47, 0.11, 0.11, 0.11)]
        fonts = Path(__file__).with_name('fonts')
        self.add_font('DejaVu', fname=fonts / 'DejaVuSans.ttf')
        self.add_font('DejaVu', style='B',
            fname=fonts / 'DejaVuSans-Bold.ttf')
        self.alias_nb_pages()

    def header(self):
        self.checker.check()
        self.set_font('DejaVu', 'B', 11)
        self.multi_cell(0, 5, self.heading, new_x='LMARGIN', new_y='NEXT')
        self.set_font('DejaVu', size=7)
        for text in self.filters:
            self.multi_cell(0, 3.5, text, new_x='LMARGIN', new_y='NEXT')
        self.ln(2)
        self.set_font('DejaVu', 'B', 7)
        self.row(self.columns, bold=True, fill=True, paginate=False)
        self.body_top = self.y

    def footer(self):
        self.set_y(-8)
        self.set_font('DejaVu', size=7)
        self.cell(0, 4, f'{self.page_no()} / {{nb}}', align='R')

    def row(self, values, bold=False, fill=False, widths=None,
            aligns=None, bottom=False, paginate=True):
        self.checker.check()
        self.set_font('DejaVu', 'B' if bold else '', 7)
        widths = widths or self.widths
        aligns = aligns or ['L', 'L', 'L', 'R', 'R', 'R']
        lines = []
        for width, value in zip(widths, values):
            text = str(value or '')
            if '\n' not in text and self.get_string_width(text) <= width - 2:
                lines.append([text])
            else:
                lines.append(self.multi_cell(width, 4, text,
                    dry_run=True, output='LINES'))
        count = max(map(len, lines), default=1)
        # Keep normal rows together; exceptionally long descriptions may span
        # pages without losing text or repeating monetary amounts.
        if (paginate and self.y + count * 4 > self.h - self.b_margin
                and count * 4 <= self.h - self.b_margin - self.body_top):
            self.add_page()
            self.set_font('DejaVu', 'B' if bold else '', 7)
        for index in range(count):
            if paginate and self.y + 4 > self.h - self.b_margin:
                self.add_page()
                self.set_font('DejaVu', 'B' if bold else '', 7)
            self.set_x(self.l_margin)
            self.set_fill_color(235, 235, 235)
            for width, cell_lines, align in zip(widths, lines, aligns):
                text = cell_lines[index] if index < len(cell_lines) else ''
                self.cell(width, 4, text, align=align, fill=fill)
            self.ln(4)
        if bottom:
            self.line(self.l_margin, self.y, self.w - self.r_margin, self.y)


class GeneralLedgerXlsxReport(XlsxReport, metaclass=PoolMeta):
    __name__ = 'account_reports.general_ledger_xlsx'

    @classmethod
    def get_content(cls, ids, data):
        pool = Pool()
        Config = pool.get('account.configuration')

        config = Config(1)
        timeout = data.get('timeout') or config.default_timeout or 300
        checker = TimeoutChecker(timeout, GeneralLedgerReport.timeout_exception)

        start_prepare = datetime.now()
        with Transaction().set_context(active_test=False):
            try:
                records, parameters = GeneralLedgerReport.prepare(data, checker)
            except TimeoutException:
                raise UserError(gettext('account_reports.msg_timeout_exception'))
        end_prepare = datetime.now()

        context = cls._xlsx_context()
        if timeout:
            context['timeout_report'] = (
                timeout - int((end_prepare - start_prepare).total_seconds()))

        with Transaction().set_context(**context):
            return cls._build_workbook(records, parameters)

    @classmethod
    def _build_workbook(cls, records, parameters):
        wb = Workbook()
        ws = wb.active
        ws.title = _('General Ledger')[:31]

        def xls(value, **kwargs):
            return convert_str_to_float(html_render(value, **kwargs))

        ws.append([parameters['company'], _('General Ledger'),
            html_render(datetime.now())])
        ws.append(['%s: %s' % (
            parameters['company_vat_label'], parameters['company_vat'])])
        if parameters['start_date']:
            ws.append([_('Start Date: %s End Date: %s') % (
                parameters['start_date'], parameters['end_date'])])
        else:
            start_period = (parameters['start_period'].name
                if parameters['start_period'] else '')
            end_period = (parameters['end_period'].name
                if parameters['end_period'] else '')
            ws.append([_('Fiscal Year: %s Start Period: %s End Period: %s')
                % (parameters['fiscal_year'], start_period, end_period)])
        if parameters['parties']:
            ws.append([_('Parties: %s') % parameters['parties']])
        else:
            ws.append([_('All Parties')])
        if parameters['accounts']:
            ws.append([_('Accounts: %s') % parameters['accounts']])
        else:
            ws.append([_('All Accounts')])
        ws.append([_(
            "When move number is between parentheses it means that "
            "it has no post number and the number shown is the "
            "provisional one.")])
        ws.append([])

        ws.append([
            _('Date'),
            _('Number'),
            _('Reference // Description'),
            _('Debit'),
            _('Credit'),
            _('Balance'),
            ])

        show_description = parameters.get('show_description', True)
        for record in records.values():
            ws.append([
                record['code'],
                '',
                record['party'] or record['account'],
                _('Previous balance...') if record['lines'] else '',
                '',
                xls(record['previous_balance']) if record['lines'] else '',
                ])

            if record['lines']:
                for line_info in record['lines']:
                    if line_info['line']:
                        line = line_info['line']
                        if line.move and line.move.number:
                            number = line.move.number
                        else:
                            number = line.party and line.party.accounting_name or ''
                        description = ''
                        if line_info['ref']:
                            description += line_info['ref']
                        if (line_info['ref'] and show_description
                                and line and (line.description
                                    or line.move_description_used)):
                            description += ' // '
                        if show_description and line and line.description:
                            description += ' %s ' % line.description
                        elif (show_description and line
                                and line.move_description_used):
                            description += ' %s ' % line.move_description_used
                        ws.append([
                            html_render(line.date),
                            number,
                            description,
                            xls(line_info['debit']),
                            xls(line_info['credit']),
                            xls(line_info['balance']),
                            ])
                    else:
                        ws.append([
                            '',
                            '-',
                            _('Previous balance'),
                            xls(line_info['debit']),
                            xls(line_info['credit']),
                            xls(line_info['balance']),
                            ])
            else:
                ws.append([
                    '',
                    '-',
                    _('Previous balance'),
                    xls(record['total_debit']),
                    xls(record['total_credit']),
                    xls(record['total_debit'] - record['total_credit']),
                    ])

            ws.append([
                '',
                '',
                _('Total Fiscal Year'),
                xls(record['total_debit']),
                xls(record['total_credit']),
                xls(record['total_debit'] - record['total_credit']),
                ])
            ws.append([
                record['code'],
                '',
                record['party'] if record['party'] else record['account'],
                _('Total'),
                '',
                xls(record['previous_balance']
                    + record['total_debit']
                    - record['total_credit']),
                ])
            ws.append([])

        return save_workbook(wb)
