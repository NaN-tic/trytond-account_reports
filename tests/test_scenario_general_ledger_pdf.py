import datetime
import tempfile
import time
import unittest
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from proteus import Model
from pypdf import PdfReader

from trytond.modules.account.tests.tools import (
    create_chart, create_fiscalyear, get_accounts)
from trytond.modules.account_invoice.tests.tools import (
    set_fiscalyear_invoice_sequences)
from trytond.modules.company.tests.tools import create_company, get_company
from trytond.pool import Pool
from trytond.tests.test_tryton import drop_db
from trytond.tests.tools import activate_modules
from trytond.transaction import Transaction


class TestGeneralLedgerPDF(unittest.TestCase):

    def setUp(self):
        drop_db()
        super().setUp()

    def tearDown(self):
        drop_db()
        super().tearDown()

    def test(self):
        config = activate_modules('account_reports')
        create_company(config=config)
        company = get_company(config=config)
        create_chart(company, config=config)
        fiscalyear = set_fiscalyear_invoice_sequences(create_fiscalyear(
            company, today=(datetime.date(2025, 1, 1),
                datetime.date(2025, 12, 31)), config=config), config=config)
        fiscalyear.click('create_period')
        first, second = fiscalyear.periods[:2]
        accounts = get_accounts(company, config=config)
        accounts['expense'].reconcile = True
        accounts['expense'].save()
        Move = Model.get('account.move', config=config)
        Journal = Model.get('account.journal', config=config)
        journal, = Journal.find([('code', '=', 'EXP')])
        for period, count, amount in [(first, 1, Decimal('12.34')),
                (second, 80, Decimal('10.00'))]:
            move = Move(company=company, period=period, journal=journal,
                date=period.start_date)
            for index in range(count):
                line = move.lines.new()
                line.account = accounts['expense']
                line.debit = amount
                line.description = f'Apunt-{index:03d} català € àéïòü'
                if index == 40:
                    line.description += ' ' + 'Descripció llarga ' * 500
                    line.description += 'FINAL-DESCRIPCIO'
            line = move.lines.new()
            line.account = accounts['cash']
            line.credit = count * amount
            move.save()
            move.click('post')

        with Transaction().start(config.database_name, 0,
                context={'company': company.id, 'language': 'en'}):
            Report = Pool().get('account_reports.general_ledger', type='report')
            data = dict(company=company.id, fiscalyear=fiscalyear.id,
                start_period=second.id, end_period=second.id,
                start_date=None, end_date=None,
                accounts=[accounts['expense'].id], parties=[],
                all_accounts=True, show_description=True,
                output_format='pdf', timeout=300)
            started = time.monotonic()
            extension, content, _, filename = Report.execute([], data)
            print('Ledger PDF seconds:', time.monotonic() - started)
            self.assertEqual(extension, 'pdf')
            self.assertTrue(filename)
            reader = PdfReader(BytesIO(content))
            self.assertGreater(len(reader.pages), 2)
            text = '\n'.join(page.extract_text() for page in reader.pages)
            for index in range(80):
                self.assertIn(f'Apunt-{index:03d}', text)
            for expected in ['català € àéïòü', 'FINAL-DESCRIPCIO',
                    '12.34', '800.00', '812.34', 'Total Fiscal Year']:
                self.assertIn(expected, text)
            for index, page in enumerate(reader.pages, 1):
                page_text = page.extract_text()
                self.assertIn('General Ledger', page_text)
                self.assertIn('Reference // Description', page_text)
                self.assertIn(f'{index} / {len(reader.pages)}', page_text)
            with tempfile.NamedTemporaryFile(prefix='general-ledger-',
                    suffix='.pdf', delete=False) as output:
                output.write(content)
                print('Ledger PDF preview:', output.name)

            Start = Pool().get('account_reports.print_general_ledger.start')
            self.assertEqual(dict(Start.output_format.selection),
                {'pdf': 'PDF', 'xlsx': 'Excel'})
            Xlsx = Pool().get('account_reports.general_ledger_xlsx',
                type='report')
            workbook = load_workbook(BytesIO(Xlsx.get_content([], data)))
            self.assertIn(812.34, [cell.value
                for row in workbook.active for cell in row])

            data.update(output_format='pdf', start_period=None, end_period=None,
                start_date=datetime.date(2025, 3, 1),
                end_date=datetime.date(2025, 3, 31), show_description=False)
            _, content, _, _ = Report.execute([], data)
            text = '\n'.join(page.extract_text()
                for page in PdfReader(BytesIO(content)).pages)
            self.assertIn('Start Date:', text)
            self.assertIn('Previous balance', text)
            self.assertIn('812.34', text)
            self.assertNotIn('Apunt-', text)


            # The same posted entries feed the journals, open lines and
            # comparative trial balance, including their PDF pagination.
            journal_data = dict(company=company.id, fiscalyear=fiscalyear.id,
                start_period=first.id, end_period=second.id, journals=[],
                open_close_account_moves=False, open_move_description='Open',
                close_move_description='Close', output_format='pdf')
            trial_data = dict(company=company.id, fiscalyear=fiscalyear.id,
                start_period=second.id, end_period=second.id,
                start_date=None, end_date=None, comparison_fiscalyear=fiscalyear.id,
                comparison_start_period=first.id, comparison_end_period=first.id,
                comparison_start_date=None, comparison_end_date=None,
                digits=None, add_initial_balance=True, only_moves=False,
                moves_or_initial=False, show_summary_lines=True, split_parties=False,
                accounts=[accounts['expense'].id], parties=[], output_format='pdf',
                timeout=300)
            cases = [
                ('journal', journal_data, ['Total month', '812.34', 'Apunt-079']),
                ('abreviated_journal', dict(company=company.id,
                    fiscalyear=fiscalyear.id, display_account='bal_movement', level=6,
                    output_format='pdf'), ['Total Period', '812.34']),
                ('open_move_lines', dict(company=company.id,
                    date=second.end_date, accounts=[accounts['expense'].id],
                    parties=[], show_description=True, output_format='pdf',
                    timeout=300), ['812.34', 'Apunt-079']),
                ('trial_balance', trial_data,
                    ['Main Balance', 'Comparision Balance', '800.00', '12.34']),
                ]
            for name, report_data, expected in cases:
                with self.subTest(report=name):
                    Report = Pool().get('account_reports.' + name, type='report')
                    extension, content, _, _ = Report.execute([], report_data)
                    self.assertEqual(extension, 'pdf')
                    reader = PdfReader(BytesIO(content))
                    text = '\n'.join(page.extract_text() for page in reader.pages)
                    for value in expected:
                        self.assertIn(value, text)
                    if name != 'trial_balance':
                        self.assertGreater(len(reader.pages), 1)
                    for index, page in enumerate(reader.pages, 1):
                        self.assertIn(f'{index} / {len(reader.pages)}', page.extract_text())
                    with tempfile.NamedTemporaryFile(prefix=name + '-',
                            suffix='.pdf', delete=False) as output:
                        output.write(content)
                        print('Report PDF preview:', output.name)
                    Start = Pool().get('account_reports.print_' + name + '.start')
                    self.assertNotIn('html', dict(Start.output_format.selection))
                    Xlsx = Pool().get('account_reports.' + name + '_xlsx', type='report')
                    workbook = load_workbook(BytesIO(Xlsx.get_content([], report_data)))
                    self.assertGreater(workbook.active.max_row, 5)
