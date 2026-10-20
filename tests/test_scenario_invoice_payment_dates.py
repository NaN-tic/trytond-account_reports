import datetime
import unittest
import tempfile
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader
from decimal import Decimal

from dateutil.relativedelta import relativedelta
from proteus import Model

from trytond.modules.account.tests.tools import (
    create_chart,
    create_tax,
    create_fiscalyear,
    get_accounts,
    )
from trytond.modules.account_invoice.tests.tools import (
    set_fiscalyear_invoice_sequences,
    )
from trytond.modules.company.tests.tools import create_company, get_company
from trytond.pool import Pool
from trytond.tests.test_tryton import drop_db
from trytond.tests.tools import activate_modules
from trytond.transaction import Transaction


class TestInvoicePaymentDates(unittest.TestCase):

    def setUp(self):
        drop_db()
        self.config = activate_modules('account_reports')
        super().setUp()

    def tearDown(self):
        drop_db()
        super().tearDown()

    def test(self):
        Journal = Model.get('account.journal', config=self.config)
        Move = Model.get('account.move', config=self.config)
        Invoice = Model.get('account.invoice', config=self.config)
        Party = Model.get('party.party', config=self.config)
        PaymentTerm = Model.get('account.invoice.payment_term', config=self.config)

        _ = create_company(config=self.config)
        company = get_company(config=self.config)
        create_chart(company, config=self.config)
        fiscalyear = set_fiscalyear_invoice_sequences(
            create_fiscalyear(company, config=self.config),
            config=self.config)
        fiscalyear.click('create_period')

        today = datetime.date.today()
        period = next(
            period for period in fiscalyear.periods
            if period.start_date <= today <= period.end_date)

        accounts = get_accounts(company, config=self.config)
        receivable = accounts['receivable']
        payable = accounts['payable']
        revenue = accounts['revenue']
        expense = accounts['expense']
        cash = accounts['cash']

        customer = Party(name='Customer')
        customer.addresses.new()
        customer.save()
        supplier = Party(name='Supplier')
        supplier.addresses.new()
        supplier.save()
        customer_name = customer.rec_name
        supplier_name = supplier.rec_name

        payment_term = PaymentTerm(name='15 days')
        line = payment_term.lines.new(type='remainder')
        line.relativedeltas.new(days=15)
        payment_term.save()
        payment_term_name = payment_term.rec_name

        customer_invoice = Invoice()
        customer_invoice.company = company
        customer_invoice.currency = company.currency
        customer_invoice.account = receivable
        customer_invoice.type = 'out'
        customer_invoice.party = customer
        customer_invoice.invoice_date = today
        customer_invoice.payment_term = payment_term
        customer_invoice.payment_term_date = today
        customer_line = customer_invoice.lines.new()
        customer_line.account = revenue
        customer_line.quantity = 1
        customer_line.unit_price = Decimal('100')
        customer_invoice.click('post')

        supplier_invoice = Invoice()
        supplier_invoice.company = company
        supplier_invoice.currency = company.currency
        supplier_invoice.account = payable
        supplier_invoice.type = 'in'
        supplier_invoice.party = supplier
        supplier_invoice.invoice_date = today
        supplier_invoice.payment_term = payment_term
        supplier_invoice.payment_term_date = today
        supplier_line = supplier_invoice.lines.new()
        supplier_line.account = expense
        supplier_line.quantity = 1
        supplier_line.unit_price = Decimal('75')
        supplier_invoice.click('post')

        customer_invoice = Invoice.find([
                ('party', '=', customer.id),
                ('type', '=', 'out'),
                ('invoice_date', '=', today),
                ], limit=1)[0]
        supplier_invoice = Invoice.find([
                ('party', '=', supplier.id),
                ('type', '=', 'in'),
                ('invoice_date', '=', today),
                ], limit=1)[0]

        journal_cash, = Journal.find([
                ('code', '=', 'CASH'),
                ])

        customer_payment_date = today + relativedelta(days=20)
        customer_payment = Move()
        customer_payment.company = company
        customer_payment.period = next(p for p in fiscalyear.periods
            if p.start_date <= customer_payment_date <= p.end_date)
        customer_payment.journal = journal_cash
        customer_payment.date = customer_payment_date
        line = customer_payment.lines.new()
        line.account = cash
        line.debit = Decimal('100')
        line = customer_payment.lines.new()
        line.account = receivable
        line.party = customer
        line.credit = Decimal('100')
        customer_payment.save()
        customer_payment.click('post')
        customer_receivable = next(
                line for line in customer_payment.lines if line.account == receivable)

        supplier_payment_date = today + relativedelta(days=25)
        supplier_payment = Move()
        supplier_payment.company = company
        supplier_payment.period = next(p for p in fiscalyear.periods
            if p.start_date <= supplier_payment_date <= p.end_date)
        supplier_payment.journal = journal_cash
        supplier_payment.date = supplier_payment_date
        line = supplier_payment.lines.new()
        line.account = cash
        line.credit = Decimal('75')
        line = supplier_payment.lines.new()
        line.account = payable
        line.party = supplier
        line.debit = Decimal('75')
        supplier_payment.save()
        supplier_payment.click('post')
        supplier_payable = next(
                line for line in supplier_payment.lines if line.account == payable)

        with Transaction().start(self.config.database_name, 0):
            Company = Pool().get('company.company')
            FiscalYear = Pool().get('account.fiscalyear')
            Period = Pool().get('account.period')
            Line = Pool().get('account.move.line')
            Invoice = Pool().get('account.invoice')
            PrintInvoicePaymentDates = Pool().get(
                'account_reports.print_invoice_payment_dates', type='wizard')
            InvoicePaymentDatesReport = Pool().get(
                'account_reports.invoice_payment_dates', type='report')
            from trytond.modules.account_reports.common import TimeoutChecker
            company = Company.browse([company.id])[0]
            fiscalyear = FiscalYear.browse([fiscalyear.id])[0]
            period = Period.browse([period.id])[0]
            customer_invoice = Invoice.browse([customer_invoice.id])[0]
            supplier_invoice = Invoice.browse([supplier_invoice.id])[0]
            customer_invoice_line = next(iter(customer_invoice.lines_to_pay))
            supplier_invoice_line = next(iter(supplier_invoice.lines_to_pay))
            Line.reconcile([
                    Line.browse([customer_invoice_line.id])[0],
                    Line.browse([customer_receivable.id])[0],
                    ])
            Line.reconcile([
                    Line.browse([supplier_invoice_line.id])[0],
                    Line.browse([supplier_payable.id])[0],
                    ])

            session_id, _, _ = PrintInvoicePaymentDates.create()
            print_invoice_payment_dates = PrintInvoicePaymentDates(session_id)
            print_invoice_payment_dates.start.company = company
            print_invoice_payment_dates.start.fiscalyear = fiscalyear
            print_invoice_payment_dates.start.periods = [period.id]
            print_invoice_payment_dates.start.start_date = None
            print_invoice_payment_dates.start.end_date = None
            print_invoice_payment_dates.start.invoice_type = 'out'
            print_invoice_payment_dates.start.output_format = 'pdf'
            print_invoice_payment_dates.start.timeout = 30
            _, data = print_invoice_payment_dates.do_print_(None)
            checker = TimeoutChecker(
                print_invoice_payment_dates.start.timeout,
                InvoicePaymentDatesReport.timeout_exception)
            records, parameters = InvoicePaymentDatesReport.prepare(
                data, checker)

            self.assertEqual(data['company'], company.id)
            self.assertEqual(data['invoice_type'], 'out')
            self.assertEqual(parameters['company'], company.rec_name)
            self.assertEqual(parameters['invoice_type'], 'Customer Invoices')
            self.assertEqual(parameters['periods'], period.rec_name)
            self.assertEqual(len(records), 1)
            record = records[0]
            self.assertEqual(record['number'], customer_invoice.number)
            self.assertEqual(record['party'], customer_name)
            self.assertEqual(record['invoice_date'], today.strftime('%d/%m/%Y'))
            self.assertEqual(record['state'], customer_invoice.state)
            self.assertEqual(record['payment_type'], payment_term_name)
            self.assertEqual(record['due_date'],
                (today + relativedelta(days=15)).strftime('%d/%m/%Y'))
            self.assertEqual(record['payment_date'],
                customer_payment_date.strftime('%d/%m/%Y'))
            self.assertEqual(record['payment_days'], 20)
            self.assertEqual(record['total_amount'], customer_invoice.total_amount)
            self.assertEqual(record['due_amount'], customer_invoice.total_amount)

            session_id, _, _ = PrintInvoicePaymentDates.create()
            print_invoice_payment_dates = PrintInvoicePaymentDates(session_id)
            print_invoice_payment_dates.start.company = company
            print_invoice_payment_dates.start.fiscalyear = fiscalyear
            print_invoice_payment_dates.start.periods = [period.id]
            print_invoice_payment_dates.start.start_date = None
            print_invoice_payment_dates.start.end_date = None
            print_invoice_payment_dates.start.invoice_type = 'in'
            print_invoice_payment_dates.start.output_format = 'pdf'
            print_invoice_payment_dates.start.timeout = 30
            _, data = print_invoice_payment_dates.do_print_(None)
            checker = TimeoutChecker(
                print_invoice_payment_dates.start.timeout,
                InvoicePaymentDatesReport.timeout_exception)
            records, parameters = InvoicePaymentDatesReport.prepare(
                data, checker)

            self.assertEqual(data['invoice_type'], 'in')
            self.assertEqual(parameters['invoice_type'], 'Supplier Invoices')
            self.assertEqual(len(records), 1)
            record = records[0]
            self.assertEqual(record['number'], supplier_invoice.number)
            self.assertEqual(record['party'], supplier_name)
            self.assertEqual(record['state'], supplier_invoice.state)
            self.assertEqual(record['payment_type'], payment_term_name)
            self.assertEqual(record['due_date'],
                (today + relativedelta(days=15)).strftime('%d/%m/%Y'))
            self.assertEqual(record['payment_date'],
                supplier_payment_date.strftime('%d/%m/%Y'))
            self.assertEqual(record['payment_days'], 25)
            self.assertEqual(record['due_amount'], supplier_invoice.total_amount)


            extension, content, _, _ = InvoicePaymentDatesReport.execute([], data)
            self.assertEqual(extension, 'pdf')
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
            self.assertIn(supplier_name, text)
            self.assertIn('75.00', text)
            self.assertIn('25', text)
            with tempfile.NamedTemporaryFile(prefix='invoice_payment_dates-',
                    suffix='.pdf', delete=False) as output:
                output.write(content)
                print('Report PDF preview:', output.name)
            data['output_format'] = 'xlsx'
            extension, content, _, _ = InvoicePaymentDatesReport.execute([], data)
            self.assertEqual(extension, 'xlsx')
            workbook = load_workbook(BytesIO(content))
            values = [cell.value for row in workbook.active for cell in row]
            self.assertIn(supplier_name, values)
            self.assertIn(75, values)
            self.assertIn(25, values)

        # A further customer invoice with two taxes exercises repeated invoice
        # lines, per-tax totals and grouped period totals.
        company = get_company(config=self.config)
        Invoice = Model.get('account.invoice', config=self.config)
        tax = create_tax(Decimal('0.10'), company, config=self.config)
        tax.save()
        other_tax = create_tax(Decimal('0.05'), company, config=self.config)
        other_tax.save()
        taxed_invoice = Invoice(company=company, currency=company.currency,
            account=receivable, type='out', party=customer,
            invoice_date=today, payment_term=payment_term)
        line = taxed_invoice.lines.new()
        line.account = revenue
        line.quantity = 1
        line.unit_price = Decimal('200')
        line.taxes.extend([tax, other_tax])
        taxed_invoice.click('post')
        invoice_number = taxed_invoice.number
        with Transaction().start(self.config.database_name, 0,
                context={'company': company.id, 'language': 'en'}):
            Report = Pool().get('account_reports.taxes_by_invoice', type='report')
            data = dict(company=company.id, fiscalyear=fiscalyear.id,
                start_date=None, end_date=None, periods=[], parties=[customer.id],
                excluded_parties=[], partner_type='customers', grouping='invoice',
                tax_type='all', totals_only=False, taxes=[], output_format='pdf')
            extension, content, _, _ = Report.execute([], data)
            self.assertEqual(extension, 'pdf')
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
            self.assertIn(invoice_number, text)
            for amount in ['200.00', '20.00', '10.00', '230.00', '430.00']:
                self.assertIn(amount, text)
            self.assertEqual(text.count('230.00'), 1)
            with tempfile.NamedTemporaryFile(prefix='taxes_by_invoice-',
                    suffix='.pdf', delete=False) as output:
                output.write(content)
                print('Report PDF preview:', output.name)
            data['totals_only'] = True
            _, content, _, _ = Report.execute([], data)
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
            self.assertNotIn('230.00', text)
            self.assertIn('430.00', text)
            data.update(totals_only=False, grouping='base_tax_code')
            _, content, _, _ = Report.execute([], data)
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
            self.assertIn(invoice_number, text)
            self.assertIn('220.00', text)
            self.assertIn('210.00', text)
            Xlsx = Pool().get('account_reports.taxes_by_invoice_xlsx', type='report')
            workbook = load_workbook(BytesIO(Xlsx.get_content([], data)))
            values = [cell.value for row in workbook.active for cell in row]
            self.assertIn(220, values)
            self.assertIn(210, values)
            data.update(start_date=datetime.date(1900, 1, 1),
                end_date=datetime.date(1900, 1, 31))
            _, content, _, _ = Report.execute([], data)
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
            self.assertIn('No records found', text)
